import json
from flask import request, jsonify
from flask_login import login_required, current_user

from src.config.database import db_session
from src.models.models_db import QuizDB, QuestionDB
from .quiz import quiz_bp

# --- ADMIN QUIZZES CRUD ---

@quiz_bp.route('/api/questions/upload', methods=['POST'])
@login_required
def upload_question_image():
    if current_user.role != 'admin':
        return jsonify({"status": "error", "message": "Bạn không có quyền thực hiện hành động này."}), 403

    if 'image' not in request.files and 'file' not in request.files:
        return jsonify({"status": "error", "message": "Không tìm thấy file tải lên."}), 400
        
    file = request.files.get('image') or request.files.get('file')
    if file.filename == '':
        return jsonify({"status": "error", "message": "Tên file rỗng."}), 400
        
    allowed_extensions = {
        # Images
        'png', 'jpg', 'jpeg', 'gif', 'webp',
        # Videos
        'mp4', 'webm', 'ogg', 'avi', 'mov', 'mkv',
        # Audio
        'mp3', 'wav', 'm4a', 'aac'
    }
    filename = file.filename
    if not ('.' in filename and filename.rsplit('.', 1)[1].lower() in allowed_extensions):
        return jsonify({"status": "error", "message": "Định dạng file không được hỗ trợ. Vui lòng tải lên ảnh, video hoặc âm thanh phù hợp."}), 400
        
    # pyrefly: ignore [missing-import]
    from werkzeug.utils import secure_filename
    import time
    import os
    
    sec_filename = secure_filename(file.filename)
    ext = sec_filename.rsplit('.', 1)[1].lower()
    # Ensure a clean filename that preserves format
    new_filename = f"q_{int(time.time())}.{ext}"
    
    images_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), 'storage', 'images')
    os.makedirs(images_dir, exist_ok=True)
    file_path = os.path.join(images_dir, new_filename)
    
    file.save(file_path)
    
    db_image_url = f"/uploads/images/{new_filename}"
    return jsonify({
        "status": "success",
        "message": "Tải ảnh lên thành công!",
        "image_url": db_image_url
    })


@quiz_bp.route('/api/quizzes/<int:quiz_id>/questions/import-excel', methods=['POST'])
@login_required
def import_quiz_questions_excel(quiz_id):
    if current_user.role != 'admin':
        return jsonify({"status": "error", "message": "Bạn không có quyền thực hiện hành động này."}), 403

    quiz = db_session.query(QuizDB).filter(QuizDB.quiz_id == quiz_id).first()
    if not quiz:
        return jsonify({"status": "error", "message": "Không tìm thấy đề thi"}), 404

    if 'file' not in request.files:
        return jsonify({"status": "error", "message": "Không tìm thấy file Excel tải lên."}), 400

    file = request.files['file']
    if file.filename == '':
        return jsonify({"status": "error", "message": "Tên file rỗng."}), 400

    # check extension
    filename = file.filename
    ext = filename.rsplit('.', 1)[1].lower() if '.' in filename else ''
    if ext not in ('xls', 'xlsx'):
        return jsonify({"status": "error", "message": "Chỉ hỗ trợ file định dạng .xls hoặc .xlsx"}), 400

    import os
    import time
    # pyrefly: ignore [missing-import]
    from werkzeug.utils import secure_filename

    sec_filename = secure_filename(file.filename)
    temp_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), 'storage', 'temp')
    os.makedirs(temp_dir, exist_ok=True)
    temp_path = os.path.join(temp_dir, f"import_{int(time.time())}_{sec_filename}")
    file.save(temp_path)

    try:
        questions_to_add = []
        if ext == 'xls':
            import xlrd
            workbook = xlrd.open_workbook(temp_path)
            sheet = workbook.sheet_by_index(0)
            
            for r in range(1, sheet.nrows):
                row_vals = [sheet.cell_value(r, c) for c in range(sheet.ncols)]
                questions_to_add.append(row_vals)
        else:
            import openpyxl
            workbook = openpyxl.load_workbook(temp_path, read_only=True, data_only=True)
            sheet = workbook.active
            
            # Read all rows at once using iter_rows for better performance
            rows_iter = sheet.iter_rows(min_row=2, max_col=26, values_only=True)
            for row_tuple in rows_iter:
                row_vals = [val if val is not None else '' for val in row_tuple]
                if not any(row_vals):
                    continue
                # Skip guide/instruction row (row 2 in template)
                first_val = str(row_vals[0]).strip()
                second_val = str(row_vals[1] if len(row_vals) > 1 else '').strip()
                if 'Số thứ tự' in first_val or 'Ký hiệu' in first_val or 'Ký hiệu' in second_val:
                    continue
                questions_to_add.append(row_vals)
            workbook.close()

        KNOWN_TYPES = ('TF', 'YN', 'MC', 'PO', 'MR', 'PM', 'TI', 'SA', 'NUMG', 'MG', 'MS', 'SEQ', 'RNK', 'MTF', 'HS')

        # Build bulk insert mappings for maximum performance
        bulk_mappings = []

        from .quiz import check_answer  # import check_answer helper from main quiz module

        for row in questions_to_add:
            row = list(row)
            if len(row) < 2:
                continue

            shift = 0
            val_0 = str(row[0]).strip().upper()
            val_1 = str(row[1]).strip().upper()
            if val_1 in KNOWN_TYPES:
                shift = 1

            if len(row) < shift + 18:
                row += [''] * (shift + 18 - len(row))

            q_type = str(row[shift + 0]).strip().upper()
            q_text = str(row[shift + 1]).strip()
            
            if not q_type or not q_text or q_type not in KNOWN_TYPES:
                continue

            image = str(row[shift + 2]).strip()
            
            # Options (Answer 1 to Answer 10)
            answers_cols = [str(x).strip() for x in row[shift + 5 : shift + 15]]
            answers_cols = [x for x in answers_cols if x != '']

            correct_feedback = str(row[shift + 15]).strip()
            incorrect_feedback = str(row[shift + 16]).strip()

            cf_val = correct_feedback if correct_feedback else "Chính xác! Bạn làm tốt lắm"
            icf_val = incorrect_feedback if incorrect_feedback else "Sai rồi! Bạn nhớ xem lại đáp án đúng nhé"
            explanation = f"Chính xác: {cf_val}\nKhông chính xác: {icf_val}"

            points_val = row[shift + 17]
            try:
                points = int(float(points_val)) if points_val != '' else 1
            except:
                points = 1

            mapped_type = 'single_choice'
            options = []
            correct_answer = None

            if q_type in ('TF', 'YN'):
                mapped_type = 'true_false'
                opt_texts = [x.lstrip('*') for x in answers_cols] if len(answers_cols) >= 2 else ['True', 'False']
                options = [
                    {"id": "T", "text": opt_texts[0]},
                    {"id": "F", "text": opt_texts[1]}
                ]
                correct_answer = "T"
                if len(answers_cols) >= 2:
                    if answers_cols[1].startswith('*'):
                        correct_answer = "F"

            elif q_type in ('MC', 'PO'):
                mapped_type = 'single_choice'
                for idx, ans in enumerate(answers_cols):
                    opt_id = chr(65 + idx)
                    is_correct = ans.startswith('*')
                    opt_text = ans.lstrip('*')
                    options.append({"id": opt_id, "text": opt_text})
                    if is_correct:
                        correct_answer = opt_id
                if not correct_answer and options:
                    correct_answer = options[0]["id"]

            elif q_type in ('MR', 'PM'):
                mapped_type = 'multiple_choice'
                correct_answer = []
                for idx, ans in enumerate(answers_cols):
                    opt_id = chr(65 + idx)
                    is_correct = ans.startswith('*')
                    opt_text = ans.lstrip('*')
                    options.append({"id": opt_id, "text": opt_text})
                    if is_correct:
                        correct_answer.append(opt_id)

            elif q_type in ('TI', 'SA', 'NUMG'):
                mapped_type = 'short_answer'
                valid_answers = [x.lstrip('=') for x in answers_cols]
                options = []
                correct_answer = valid_answers if len(valid_answers) > 1 else (valid_answers[0] if valid_answers else "")

            elif q_type in ('MG', 'MS'):
                mapped_type = 'matching'
                left_opts = []
                right_opts = []
                correct_answer = {}
                for idx, ans in enumerate(answers_cols):
                    if '|' in ans:
                        left_txt, right_txt = ans.split('|', 1)
                    else:
                        left_txt, right_txt = ans, ans
                    left_id = f"L{idx + 1}"
                    right_id = f"R{idx + 1}"
                    left_opts.append({"id": left_id, "text": left_txt.strip()})
                    right_opts.append({"id": right_id, "text": right_txt.strip()})
                    correct_answer[left_id] = right_id
                options = {"left": left_opts, "right": right_opts}

            elif q_type in ('SEQ', 'RNK'):
                mapped_type = 'ordering'
                for idx, ans in enumerate(answers_cols):
                    opt_id = str(idx + 1)
                    options.append({"id": opt_id, "text": ans})
                correct_answer = [x["id"] for x in options]

            elif q_type == 'MTF':
                mapped_type = 'multi_true_false'
                statements = []
                correct_answer = {}
                for idx, ans in enumerate(answers_cols):
                    stmt_id = f"S{idx + 1}"
                    is_correct = ans.startswith('*')
                    stmt_text = ans.lstrip('*').strip()
                    statements.append({"id": stmt_id, "text": stmt_text})
                    correct_answer[stmt_id] = "T" if is_correct else "F"
                options = {"statements": statements, "columnLeft": "Đúng", "columnRight": "Sai"}

            elif q_type == 'HS':
                mapped_type = 'hotspot'
                all_zones = []
                correct_zones = []
                for ans in answers_cols:
                    is_correct = ans.startswith('*')
                    clean_ans = ans.lstrip('*').strip()
                    try:
                        parts = [float(v.strip()) for v in clean_ans.split(',')]
                        if len(parts) >= 4:
                            zone = {
                                "x": parts[0],
                                "y": parts[1],
                                "w": parts[2],
                                "h": parts[3],
                                "width": parts[2],
                                "height": parts[3],
                                "is_correct": is_correct
                            }
                            all_zones.append(zone)
                            if is_correct:
                                correct_zones.append({
                                    "x": parts[0],
                                    "y": parts[1],
                                    "w": parts[2],
                                    "h": parts[3],
                                    "width": parts[2],
                                    "height": parts[3]
                                })
                    except:
                        pass
                options = all_zones if all_zones else []
                if len(correct_zones) == 1:
                    correct_answer = correct_zones[0]
                elif correct_zones:
                    correct_answer = correct_zones
                else:
                    if all_zones:
                        all_zones[0]["is_correct"] = True
                        correct_answer = {
                            "x": all_zones[0]["x"],
                            "y": all_zones[0]["y"],
                            "w": all_zones[0]["w"],
                            "h": all_zones[0]["h"],
                            "width": all_zones[0]["width"],
                            "height": all_zones[0]["height"]
                        }
                    else:
                        correct_answer = {"x": 10.0, "y": 10.0, "w": 10.0, "h": 10.0, "width": 10.0, "height": 10.0}

            else:
                mapped_type = 'single_choice'
                for idx, ans in enumerate(answers_cols):
                    opt_id = chr(65 + idx)
                    options.append({"id": opt_id, "text": ans.lstrip('*')})
                    if ans.startswith('*'):
                        correct_answer = opt_id
                if not correct_answer and options:
                    correct_answer = options[0]["id"]

            # Append mapping dict for bulk insert (no ORM object creation overhead)
            bulk_mappings.append({
                'quiz_id': quiz_id,
                'question_type': mapped_type,
                'question_text': q_text,
                'image_url': image if image else None,
                'options_json': json.dumps(options, ensure_ascii=False),
                'correct_answer_json': json.dumps(correct_answer, ensure_ascii=False),
                'explanation': explanation,
                'points': points
            })

        # Bulk insert all questions in a single operation
        if bulk_mappings:
            db_session.bulk_insert_mappings(QuestionDB, bulk_mappings)
        db_session.commit()

        imported_count = len(bulk_mappings)

        try:
            os.remove(temp_path)
        except:
            pass

        return jsonify({
            "status": "success",
            "message": f"Nhập thành công {imported_count} câu hỏi từ file Excel."
        })

    except Exception as e:
        try:
            os.remove(temp_path)
        except:
            pass
        return jsonify({"status": "error", "message": f"Lỗi khi xử lý file Excel: {str(e)}"}), 500


@quiz_bp.route('/api/quizzes', methods=['GET'])
@login_required
def get_quizzes():
    quizzes = db_session.query(QuizDB).order_by(QuizDB.created_at.desc()).all()
    res = []
    for q in quizzes:
        res.append({
            "quiz_id": q.quiz_id,
            "title": q.title,
            "description": q.description or "",
            "time_limit": q.time_limit,
            "is_published": q.is_published,
            "school_level": q.school_level or "",
            "grade_level": q.grade_level,
            "quiz_mode": q.quiz_mode or "testing",
            "allow_practice": bool(q.allow_practice) if q.allow_practice is not None else False,
            "created_at": q.created_at.strftime('%Y-%m-%d %H:%M:%S') if q.created_at else "",
            "question_count": len(q.questions)
        })
    return jsonify(res)


@quiz_bp.route('/api/quizzes/<int:quiz_id>', methods=['GET'])
@login_required
def get_quiz_detail(quiz_id):
    quiz = db_session.query(QuizDB).filter(QuizDB.quiz_id == quiz_id).first()
    if not quiz:
        return jsonify({"status": "error", "message": "Không tìm thấy đề thi"}), 404
        
    can_see_answers = current_user.role == 'admin' or current_user.has_permission('trac_nghiem_ly_thuyet')
    questions = []
    for q in quiz.questions:
        try:
            options = json.loads(q.options_json)
        except Exception:
            options = q.options_json
            
        try:
            correct = json.loads(q.correct_answer_json)
        except Exception:
            correct = q.correct_answer_json
            
        questions.append({
            "question_id": q.question_id,
            "question_type": q.question_type,
            "question_text": q.question_text,
            "image_url": q.image_url or "",
            "options": options,
            "correct_answer": correct if can_see_answers else None,
            "explanation": (q.explanation or "") if can_see_answers else "",
            "points": q.points
        })
        
    return jsonify({
        "quiz_id": quiz.quiz_id,
        "title": quiz.title,
        "description": quiz.description or "",
        "time_limit": quiz.time_limit,
        "is_published": quiz.is_published,
        "school_level": quiz.school_level or "",
        "grade_level": quiz.grade_level,
        "quiz_mode": quiz.quiz_mode or "testing",
        "allow_practice": quiz.allow_practice,
        "created_at": quiz.created_at.strftime('%Y-%m-%d %H:%M:%S') if quiz.created_at else "",
        "questions": questions
    })


@quiz_bp.route('/api/quizzes', methods=['POST'])
@login_required
def create_quiz():
    if current_user.role != 'admin':
        return jsonify({"status": "error", "message": "Bạn không có quyền quản trị."}), 403
        
    data = request.get_json() or {}
    title = data.get('title', '').strip()
    description = data.get('description', '').strip()
    time_limit = int(data.get('time_limit', 45))
    is_published = bool(data.get('is_published', False))
    school_level = data.get('school_level', '').strip()
    grade_level = data.get('grade_level')
    
    if grade_level is not None and str(grade_level).strip() != '':
        try:
            grade_level = int(grade_level)
        except ValueError:
            grade_level = None
    else:
        grade_level = None
    
    quiz_mode = data.get('quiz_mode', 'testing').strip()
    if quiz_mode not in ('training', 'testing', 'practice'):
        quiz_mode = 'testing'
        
    allow_practice = bool(data.get('allow_practice', False))
    
    if not title:
        return jsonify({"status": "error", "message": "Tiêu đề không được để trống"}), 400
        
    new_quiz = QuizDB(
        title=title,
        description=description if description else None,
        time_limit=time_limit,
        is_published=is_published,
        school_level=school_level if school_level else None,
        grade_level=grade_level,
        quiz_mode=quiz_mode,
        allow_practice=allow_practice
    )
    db_session.add(new_quiz)
    db_session.commit()
    
    return jsonify({
        "status": "success",
        "message": "Đã tạo đề thi thành công",
        "quiz_id": new_quiz.quiz_id
    })


@quiz_bp.route('/api/quizzes/generate-mixed', methods=['POST'])
@login_required
def generate_mixed_quiz():
    if current_user.role != 'admin':
        return jsonify({"status": "error", "message": "Bạn không có quyền quản trị."}), 403
        
    data = request.get_json() or {}
    title = data.get('title', '').strip()
    description = data.get('description', '').strip()
    time_limit = int(data.get('time_limit', 45))
    school_level = data.get('school_level', '').strip()
    grade_level = data.get('grade_level')
    sources = data.get('sources', []) # list of {"quiz_id": int/str, "count": int}
    
    if grade_level is not None and str(grade_level).strip() != '':
        try:
            grade_level = int(grade_level)
        except ValueError:
            grade_level = None
    else:
        grade_level = None
        
    if not title:
        return jsonify({"status": "error", "message": "Tiêu đề không được để trống"}), 400
        
    # 1. Create new quiz
    new_quiz = QuizDB(
        title=title,
        description=description if description else None,
        time_limit=time_limit,
        is_published=False,
        school_level=school_level if school_level else None,
        grade_level=grade_level,
        quiz_mode='testing'
    )
    db_session.add(new_quiz)
    db_session.flush() # get new_quiz.quiz_id
    
    import random
    import json
    # 2. Extract random questions from sources
    total_copied = 0
    for src in sources:
        src_quiz_id = src.get('quiz_id')
        if not src_quiz_id:
            continue
        try:
            src_quiz_id = int(src_quiz_id)
        except ValueError:
            continue
            
        count = int(src.get('count', 0))
        if count <= 0:
            continue
            
        questions = db_session.query(QuestionDB).filter(QuestionDB.quiz_id == src_quiz_id).all()
        if not questions:
            continue
            
        # Select random questions
        selected_count = min(count, len(questions))
        selected_questions = random.sample(questions, selected_count)
        
        for q in selected_questions:
            cloned_q = QuestionDB(
                quiz_id=new_quiz.quiz_id,
                question_type=q.question_type,
                question_text=q.question_text,
                image_url=q.image_url,
                options_json=q.options_json,
                correct_answer_json=q.correct_answer_json,
                points=q.points,
                explanation=q.explanation
            )
            db_session.add(cloned_q)
            total_copied += 1
            
    db_session.commit()
    
    return jsonify({
        "status": "success",
        "message": f"Đã trộn và tạo đề thi mới thành công với {total_copied} câu hỏi.",
        "quiz_id": new_quiz.quiz_id
    })


@quiz_bp.route('/api/quizzes/<int:quiz_id>', methods=['PUT'])
@login_required
def update_quiz(quiz_id):
    if not current_user.has_permission('trac_nghiem_ly_thuyet'):
        return jsonify({"status": "error", "message": "Bạn không có quyền thực hiện hành động này."}), 403
        
    quiz = db_session.query(QuizDB).filter(QuizDB.quiz_id == quiz_id).first()
    if not quiz:
        return jsonify({"status": "error", "message": "Không tìm thấy đề thi"}), 404
        
    data = request.get_json() or {}
    title = data.get('title')
    description = data.get('description')
    time_limit = data.get('time_limit')
    is_published = data.get('is_published')
    school_level = data.get('school_level')
    grade_level = data.get('grade_level')
    allow_practice = data.get('allow_practice')
    
    if title is not None:
        quiz.title = title.strip()
    if description is not None:
        quiz.description = description.strip() if description.strip() else None
    if time_limit is not None:
        quiz.time_limit = int(time_limit)
    if is_published is not None:
        quiz.is_published = bool(is_published)
    if school_level is not None:
        quiz.school_level = school_level.strip() if school_level.strip() else None
    if grade_level is not None:
        if str(grade_level).strip() != '':
            try:
                quiz.grade_level = int(grade_level)
            except ValueError:
                quiz.grade_level = None
        else:
            quiz.grade_level = None
    if allow_practice is not None:
        quiz.allow_practice = bool(allow_practice)
            
    quiz_mode = data.get('quiz_mode')
    if quiz_mode is not None:
        if quiz_mode in ('training', 'testing', 'practice'):
            quiz.quiz_mode = quiz_mode
        
    db_session.commit()
    return jsonify({"status": "success", "message": "Cập nhật đề thi thành công"})


@quiz_bp.route('/api/quizzes/batch-practice', methods=['POST'])
@login_required
def batch_update_practice():
    if not current_user.has_permission('trac_nghiem_ly_thuyet'):
        return jsonify({"status": "error", "message": "Bạn không có quyền thực hiện hành động này."}), 403
        
    data = request.get_json() or {}
    grade_level = data.get('grade_level')
    allow_practice = data.get('allow_practice')
    
    if grade_level is None or allow_practice is None:
        return jsonify({"status": "error", "message": "Thiếu thông tin grade_level hoặc allow_practice"}), 400
    
    try:
        grade_level = int(grade_level)
    except (ValueError, TypeError):
        return jsonify({"status": "error", "message": "grade_level không hợp lệ"}), 400
    
    quizzes = db_session.query(QuizDB).filter(
        QuizDB.grade_level == grade_level,
        QuizDB.is_published == True
    ).all()
    
    count = 0
    for quiz in quizzes:
        quiz.allow_practice = bool(allow_practice)
        count += 1
    
    db_session.commit()
    
    action_text = "bật" if allow_practice else "tắt"
    return jsonify({
        "status": "success", 
        "message": f"Đã {action_text} chế độ luyện tập cho {count} đề thi Khối {grade_level}."
    })


@quiz_bp.route('/api/quizzes/<int:quiz_id>', methods=['DELETE'])
@login_required
def delete_quiz(quiz_id):
    if current_user.role != 'admin':
        return jsonify({"status": "error", "message": "Bạn không có quyền quản trị."}), 403
        
    quiz = db_session.query(QuizDB).filter(QuizDB.quiz_id == quiz_id).first()
    if not quiz:
        return jsonify({"status": "error", "message": "Không tìm thấy đề thi"}), 404
        
    db_session.delete(quiz)
    db_session.commit()
    return jsonify({"status": "success", "message": "Xóa đề thi thành công"})


# --- ADMIN QUESTIONS CRUD ---

@quiz_bp.route('/api/quizzes/<int:quiz_id>/questions', methods=['POST'])
@login_required
def create_question(quiz_id):
    if current_user.role != 'admin':
        return jsonify({"status": "error", "message": "Bạn không có quyền quản trị."}), 403
        
    quiz = db_session.query(QuizDB).filter(QuizDB.quiz_id == quiz_id).first()
    if not quiz:
        return jsonify({"status": "error", "message": "Không tìm thấy đề thi"}), 404
        
    data = request.get_json() or {}
    question_type = data.get('question_type', '').strip()
    question_text = data.get('question_text', '').strip()
    image_url = data.get('image_url', '').strip()
    options = data.get('options')
    correct_answer = data.get('correct_answer')
    points = int(data.get('points', 1))
    explanation = data.get('explanation', '').strip()
    
    if not question_type:
        return jsonify({"status": "error", "message": "Loại câu hỏi không được để trống"}), 400
        
    new_q = QuestionDB(
        quiz_id=quiz_id,
        question_type=question_type,
        question_text=question_text,
        image_url=image_url if image_url else None,
        options_json=json.dumps(options, ensure_ascii=False) if options is not None else "[]",
        correct_answer_json=json.dumps(correct_answer, ensure_ascii=False) if correct_answer is not None else "null",
        points=points,
        explanation=explanation if explanation else None
    )
    db_session.add(new_q)
    db_session.commit()
    
    return jsonify({
        "status": "success",
        "message": "Thêm câu hỏi thành công",
        "question_id": new_q.question_id
    })


@quiz_bp.route('/api/questions/<int:question_id>', methods=['PUT'])
@login_required
def update_question(question_id):
    if current_user.role != 'admin':
        return jsonify({"status": "error", "message": "Bạn không có quyền quản trị."}), 403
        
    q = db_session.query(QuestionDB).filter(QuestionDB.question_id == question_id).first()
    if not q:
        return jsonify({"status": "error", "message": "Không tìm thấy câu hỏi"}), 404
        
    data = request.get_json() or {}
    question_type = data.get('question_type')
    question_text = data.get('question_text')
    image_url = data.get('image_url')
    options = data.get('options')
    correct_answer = data.get('correct_answer')
    points = data.get('points')
    explanation = data.get('explanation')
    
    if question_type is not None:
        q.question_type = question_type.strip()
    if question_text is not None:
        q.question_text = question_text.strip()
    if image_url is not None:
        q.image_url = image_url.strip() if image_url.strip() else None
    if options is not None:
        q.options_json = json.dumps(options, ensure_ascii=False)
    if correct_answer is not None:
        q.correct_answer_json = json.dumps(correct_answer, ensure_ascii=False)
    if points is not None:
        q.points = int(points)
    if explanation is not None:
        q.explanation = explanation.strip() if explanation.strip() else None
        
    db_session.commit()
    return jsonify({"status": "success", "message": "Cập nhật câu hỏi thành công"})


@quiz_bp.route('/api/questions/<int:question_id>', methods=['DELETE'])
@login_required
def delete_question(question_id):
    if current_user.role != 'admin':
        return jsonify({"status": "error", "message": "Bạn không có quyền quản trị."}), 403
        
    q = db_session.query(QuestionDB).filter(QuestionDB.question_id == question_id).first()
    if not q:
        return jsonify({"status": "error", "message": "Không tìm thấy câu hỏi"}), 404
        
    db_session.delete(q)
    db_session.commit()
    return jsonify({"status": "success", "message": "Xóa câu hỏi thành công"})


@quiz_bp.route('/api/questions/batch-delete', methods=['POST'])
@login_required
def batch_delete_questions():
    if current_user.role != 'admin':
        return jsonify({"status": "error", "message": "Bạn không có quyền quản trị."}), 403
        
    data = request.get_json() or {}
    question_ids = data.get('question_ids', [])
    if not question_ids:
        return jsonify({"status": "error", "message": "Danh sách ID câu hỏi rỗng."}), 400
        
    questions = db_session.query(QuestionDB).filter(QuestionDB.question_id.in_(question_ids)).all()
    deleted_count = 0
    for q in questions:
        db_session.delete(q)
        deleted_count += 1
    db_session.commit()
    return jsonify({"status": "success", "message": f"Đã xóa thành công {deleted_count} câu hỏi."})


@quiz_bp.route('/api/quizzes/<int:quiz_id>/export-standalone', methods=['POST'])
@login_required
def export_quiz_standalone(quiz_id):
    if current_user.role != 'admin' and not current_user.has_permission('trac_nghiem_ly_thuyet'):
        return jsonify({"status": "error", "message": "Bạn không có quyền thực hiện hành động này."}), 403

    quiz = db_session.query(QuizDB).filter(QuizDB.quiz_id == quiz_id).first()
    if not quiz:
        return jsonify({"status": "error", "message": "Không tìm thấy đề thi."}), 404

    # Extract custom filename from request body
    data = request.get_json() or {}
    custom_filename = data.get('filename')

    # Extract questions
    questions = []
    for q in quiz.questions:
        try:
            options = json.loads(q.options_json)
        except Exception:
            options = q.options_json
            
        try:
            correct = json.loads(q.correct_answer_json)
        except Exception:
            correct = q.correct_answer_json
            
        questions.append({
            "question_id": q.question_id,
            "question_type": q.question_type,
            "question_text": q.question_text,
            "image_url": q.image_url or "",
            "options": options,
            "correct_answer": correct,
            "points": q.points
        })

    # Export to quizzes directory
    import os
    backend_dir = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    quizzes_dir = os.path.abspath(os.path.join(backend_dir, '../quizzes'))
    
    from src.utils.quiz_exporter import export_standalone_html
    try:
        filename = export_standalone_html(quiz, questions, quizzes_dir, filename=custom_filename)
        return jsonify({
            "status": "success",
            "message": f"Xuất bản file trắc nghiệm {filename} thành công!",
            "filename": filename
        })
    except Exception as e:
        return jsonify({"status": "error", "message": f"Lỗi xuất bản: {str(e)}"}), 500

